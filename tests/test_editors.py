"""The Word document editor and the LaTeX editor tabs.

    python -m pytest tests/test_editors.py
"""
import os
import sys
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import pytest
from PySide6.QtGui import QColor, QImage, QTextCursor
from PySide6.QtWidgets import QApplication

from pdfannotator import word_io
from pdfannotator.latex import engines


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def sample_docx(tmp_path):
    from docx import Document
    from docx.shared import Inches, RGBColor

    doc = Document()
    doc.add_heading("Project Report", 0)
    doc.add_heading("Introduction", 1)
    p = doc.add_paragraph("Plain, ")
    p.add_run("bold").bold = True
    red = p.add_run(" red italic")
    red.italic = True
    red.font.color.rgb = RGBColor(200, 0, 0)
    doc.add_paragraph("First bullet", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for r in range(2):
        for c in range(3):
            table.cell(r, c).text = f"r{r}c{c}"
    table.cell(0, 0).merge(table.cell(0, 1))
    image = QImage(60, 30, QImage.Format_RGB32)
    image.fill(QColor(30, 120, 200))
    image.save(str(tmp_path / "img.png"))
    doc.add_paragraph().add_run().add_picture(str(tmp_path / "img.png"), width=Inches(1))
    doc.add_page_break()
    doc.add_paragraph("After the break.")
    path = tmp_path / "src.docx"
    doc.save(path)
    return path


def test_docx_round_trip_keeps_styles_formatting_tables_and_pictures(app, sample_docx, tmp_path):
    from docx import Document

    qdoc, page = word_io.load_docx(str(sample_docx))
    assert round(page.width_pt) == 612
    out = tmp_path / "out.docx"
    word_io.save_docx(qdoc, str(out), base=str(sample_docx))
    word = Document(out)
    styles = [p.style.name for p in word.paragraphs]
    assert styles[:5] == ["Title", "Heading 1", "Normal", "List Bullet", "List Number"]
    runs = {r.text: r for r in word.paragraphs[2].runs}
    assert runs["bold"].bold and runs[" red italic"].italic and str(runs[" red italic"].font.color.rgb) == "C80000"
    assert len(word.tables) == 1 and word.tables[0].rows[0].cells[0]._tc is word.tables[0].rows[0].cells[1]._tc
    assert any(p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip") for p in word.paragraphs)
    assert word.paragraphs[-1].paragraph_format.page_break_before
    again, _ = word_io.load_docx(str(out))
    assert again.toPlainText() == qdoc.toPlainText()


def test_word_tab_edit_and_save(app, sample_docx, tmp_path):
    from docx import Document
    from pdfannotator.word_editor import WordTab

    tab = WordTab(None, str(sample_docx))
    assert not tab.dirty and "Project Report" in tab.document.toPlainText()
    cursor = tab.page.textCursor()
    cursor.movePosition(QTextCursor.End)
    tab.page.setTextCursor(cursor)
    tab.page.insertPlainText("\nA new line")
    tab.apply_style("Heading 2")
    tab.insert_table(2, 2)
    assert tab.dirty
    out = tmp_path / "edited.docx"
    tab.write(str(out))
    word = Document(out)
    assert any(p.text == "A new line" and p.style.name == "Heading 2" for p in word.paragraphs)
    assert len(word.tables) == 2
    assert "words" in tab.status_text()
    pdf = tab.export_pdf(str(tmp_path / "out.pdf"), ask_to_open=False)
    import pymupdf as fitz

    assert fitz.open(pdf).page_count >= 2  # the page break makes a second page


def test_log_parser_finds_errors_warnings_and_bad_boxes():
    log = "\n".join([
        "./main.tex:3: Undefined control sequence.",
        "l.3 Hello \\foo",
        "LaTeX Warning: Reference `x' on page 1 undefined on input line 5.",
        "Overfull \\hbox (12.3pt too wide) in paragraph at lines 7--8",
        "! Missing $ inserted.",
        "l.9 $x^2",
    ])
    problems = engines.parse_log(log, "main.tex")
    kinds = [(p.kind, p.line) for p in problems]
    assert kinds == [("error", 3), ("error", 9), ("warning", 5), ("badbox", 7)]


def test_custom_command_placeholders():
    argv = engines.custom_command('latexmk -pdf "{file}" -outdir={dir}', r"C:\docs\my report.tex")
    assert argv == ["latexmk", "-pdf", "my report.tex", r"-outdir=C:\docs"]


@pytest.mark.skipif(not engines.find_program("pdflatex"), reason="no LaTeX installed")
def test_latex_tab_compiles_and_previews(app, tmp_path):
    from pdfannotator.latex.editor import LatexTab

    src = tmp_path / "doc.tex"
    src.write_text("\\documentclass{article}\n\\begin{document}\nHello \\undefinedthing.\n\\end{document}\n",
                   encoding="utf-8")
    tab = LatexTab(None, str(src))
    index = tab.engine_combo.findData("pdflatex")
    tab.engine_combo.setCurrentIndex(index)
    tab.compile()
    deadline = time.time() + 180
    while tab.process is not None and time.time() < deadline:
        app.processEvents()
        time.sleep(0.05)
    assert tab.process is None, "compile did not finish"
    assert any(p.kind == "error" and p.line == 3 for p in tab.problems)
    assert tab.editor.markers.get(3) is not None
    assert tab.preview.doc is not None and tab.preview.doc.page_count == 1  # nonstopmode still makes the PDF
    tab.confirm_close()


def test_latex_editor_begin_end_and_comments(app):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from pdfannotator.latex.editor import LatexTab

    tab = LatexTab(None, None, "Blank")
    editor = tab.editor
    editor.setPlainText("")
    editor.insertPlainText("\\begin{itemize}")
    QTest.keyClick(editor, Qt.Key_Return)
    assert editor.toPlainText() == "\\begin{itemize}\n  \n\\end{itemize}"
    editor.selectAll()
    editor.toggle_comment()
    assert all(line.lstrip().startswith("%") for line in editor.toPlainText().splitlines() if line.strip())
    editor.insertPlainText("\\ref{")
    tab.confirm_close = lambda: True
