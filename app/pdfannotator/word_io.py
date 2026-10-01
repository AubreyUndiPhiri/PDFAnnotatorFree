"""Read a .docx into a QTextDocument and write one back (python-docx).

Reading keeps what an editor can show: paragraphs with their styles
(headings, lists), alignment, indents and spacing; runs with bold, italic,
underline, strike-through, font, size, colour, highlight and super/
subscript; line and page breaks, tabs, hyperlinks, inline images and tables
(with merged cells and shading); underline styles, capitals and small
capitals; paragraph shading and horizontal lines; list styles (bullets,
1. / 1) / (1), a., A., i., I. and their start numbers). Formatting is resolved through the run,
paragraph and base styles, so the page looks the way Word shows it.

Writing starts from the file that was opened (when there is one) and only
replaces its body, so its styles, numbering, headers, footers and page setup
survive. The Word style of every paragraph is remembered while editing
(STYLE_PROP), as is the list numbering it came from (NUM_PROP).
"""
import io
import os
from dataclasses import dataclass

from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt, QUrl
from PySide6.QtGui import (
    QBrush, QColor, QFont, QImage, QPixmap, QTextBlockFormat, QTextCharFormat, QTextCursor, QTextDocument, QTextFormat,
    QTextImageFormat, QTextLength, QTextListFormat, QTextTable, QTextTableFormat,
)

STYLE_PROP = QTextFormat.UserProperty + 1   # the paragraph's Word style name
NUM_PROP = QTextFormat.UserProperty + 2     # "numId:ilvl" of a Word list paragraph
OUTLINE_PROP = QTextFormat.UserProperty + 3  # on a list format: "number" / "bullet" for a multilevel list
HR_PROP = QTextFormat.BlockTrailingHorizontalRulerWidth   # a horizontal line (an empty paragraph with a bottom border)
PX_PER_PT = 96 / 72
EMU_PER_PX = 9525
_VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"  # python-docx has no v: prefix
_PROPORTIONAL = QTextBlockFormat.LineHeightTypes.ProportionalHeight.value
DEFAULT_FONT = ("Calibri", 11.0)

_HEADING_SIZES = {0: 26, 1: 16, 2: 13, 3: 12, 4: 11, 5: 11, 6: 11}
_BULLETS = (QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare)
_NUMBER_STYLES = {
    "decimal": QTextListFormat.ListDecimal, "lowerLetter": QTextListFormat.ListLowerAlpha,
    "upperLetter": QTextListFormat.ListUpperAlpha, "lowerRoman": QTextListFormat.ListLowerRoman,
    "upperRoman": QTextListFormat.ListUpperRoman,
}
# bullet characters Word uses for each Qt bullet (with their fonts)
_BULLET_CHARS = {QTextListFormat.ListDisc: ("\u2022", None), QTextListFormat.ListCircle: ("o", "Courier New"),
                 QTextListFormat.ListSquare: ("\u00a7", "Wingdings")}
_NUM_FORMATS = {v: k for k, v in _NUMBER_STYLES.items()}
_UNDERLINES = {  # Word underline -> Qt underline style
    "single": QTextCharFormat.SingleUnderline, "words": QTextCharFormat.SingleUnderline,
    "double": QTextCharFormat.SingleUnderline, "thick": QTextCharFormat.SingleUnderline,
    "dotted": QTextCharFormat.DotLine, "dottedHeavy": QTextCharFormat.DotLine,
    "dash": QTextCharFormat.DashUnderline, "dashedHeavy": QTextCharFormat.DashUnderline,
    "dashLong": QTextCharFormat.DashUnderline, "dotDash": QTextCharFormat.DashDotLine,
    "dashDotHeavy": QTextCharFormat.DashDotLine, "dotDotDash": QTextCharFormat.DashDotDotLine,
    "dashDotDotHeavy": QTextCharFormat.DashDotDotLine, "wave": QTextCharFormat.WaveUnderline,
    "wavyHeavy": QTextCharFormat.WaveUnderline, "wavyDouble": QTextCharFormat.WaveUnderline,
}
_UNDERLINE_OUT = {QTextCharFormat.SingleUnderline: "single", QTextCharFormat.DotLine: "dotted",
                  QTextCharFormat.DashUnderline: "dash", QTextCharFormat.DashDotLine: "dotDash",
                  QTextCharFormat.DashDotDotLine: "dotDotDash", QTextCharFormat.WaveUnderline: "wave"}
_HIGHLIGHTS = {
    "yellow": "#ffff00", "green": "#00ff00", "cyan": "#00ffff", "magenta": "#ff00ff", "blue": "#0000ff",
    "red": "#ff0000", "darkBlue": "#000080", "darkCyan": "#008080", "darkGreen": "#008000",
    "darkMagenta": "#800080", "darkRed": "#800000", "darkYellow": "#808000", "darkGray": "#808080",
    "lightGray": "#c0c0c0", "black": "#000000", "white": "#ffffff",
}


def _qn(tag):
    from docx.oxml.ns import qn

    return qn(tag)


@dataclass
class PageSetup:
    width_pt: float = 612.0      # US Letter, python-docx's default template
    height_pt: float = 792.0
    margins_pt: tuple = (72.0, 72.0, 72.0, 72.0)   # left, top, right, bottom

    @property
    def width_px(self):
        return round(self.width_pt * PX_PER_PT)


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------

class _Reader:
    def __init__(self, word, qdoc):
        self.word = word
        self.qdoc = qdoc
        self.lists = {}          # (numId, ilvl) -> QTextList
        self.images = 0
        self.page_break = False  # the next block starts a new page
        self._numbering = self._load_numbering()
        self._defaults = self._doc_defaults()
        self._chains = {}

    # ---- styles ----------------------------------------------------------
    def _doc_defaults(self):
        family, size = DEFAULT_FONT
        try:
            rpr = self.word.styles.element.find(_qn("w:docDefaults")).find(_qn("w:rPrDefault")).find(_qn("w:rPr"))
            fonts = rpr.find(_qn("w:rFonts"))
            if fonts is not None and fonts.get(_qn("w:ascii")):
                family = fonts.get(_qn("w:ascii"))
            sz = rpr.find(_qn("w:sz"))
            if sz is not None:
                size = int(sz.get(_qn("w:val"))) / 2
        except (AttributeError, TypeError, ValueError):
            pass
        return family, size

    def _chain(self, style):
        """The style and the styles it is based on, nearest first (cached:
        python-docx looks styles up in the XML every time)."""
        if style is None:
            return []
        key = style.style_id
        if key not in self._chains:
            seen = []
            while style is not None and all(x.style_id != style.style_id for x in seen):
                seen.append(style)
                style = style.base_style
            self._chains[key] = seen
        return self._chains[key]

    @staticmethod
    def _first(chain, name):
        """The nearest style in `chain` that sets font attribute `name`."""
        for style in chain:
            try:
                value = getattr(style.font, name)
            except (AttributeError, ValueError):
                continue
            if value is not None:
                return value
        return None

    @staticmethod
    def _first_color(chain):
        for style in chain:
            try:
                if style.font.color.type is not None and style.font.color.rgb:
                    return style.font.color.rgb
            except (AttributeError, ValueError):
                continue
        return None

    def _para_attr(self, para, name, chain):
        value = getattr(para.paragraph_format, name)
        if value is not None:
            return value
        for style in chain:
            try:
                value = getattr(style.paragraph_format, name)
            except AttributeError:
                continue
            if value is not None:
                return value
        return None

    # ---- numbering -------------------------------------------------------
    def _load_numbering(self):
        """{(numId, ilvl): (numFmt, lvlText, start)} from numbering.xml."""
        try:
            root = self.word.part.numbering_part.element
        except (NotImplementedError, KeyError, AttributeError):
            return {}
        abstract = {}

        def val(lvl, tag, default):
            el = lvl.find(_qn(tag))
            return el.get(_qn("w:val")) if el is not None and el.get(_qn("w:val")) is not None else default

        for an in root.findall(_qn("w:abstractNum")):
            levels = {}
            for lvl in an.findall(_qn("w:lvl")):
                try:
                    start = int(val(lvl, "w:start", "1"))
                except ValueError:
                    start = 1
                levels[lvl.get(_qn("w:ilvl"))] = (val(lvl, "w:numFmt", "decimal"), val(lvl, "w:lvlText", ""), start)
            abstract[an.get(_qn("w:abstractNumId"))] = levels
        out = {}
        for num in root.findall(_qn("w:num")):
            ref = num.find(_qn("w:abstractNumId"))
            if ref is None:
                continue
            for ilvl, fmt in abstract.get(ref.get(_qn("w:val")), {}).items():
                out[(num.get(_qn("w:numId")), ilvl)] = fmt
        return out

    def _num_pr(self, para, chain):
        """(numId, ilvl) of a list paragraph, from the paragraph or its style."""
        sources = [para._p.pPr] + [getattr(x.element, "pPr", None) for x in chain]
        for ppr in sources:
            if ppr is None:
                continue
            num_pr = ppr.find(_qn("w:numPr"))
            if num_pr is None:
                continue
            num_id = num_pr.find(_qn("w:numId"))
            ilvl = num_pr.find(_qn("w:ilvl"))
            if num_id is None or num_id.get(_qn("w:val")) == "0":
                return None
            return num_id.get(_qn("w:val")), ilvl.get(_qn("w:val")) if ilvl is not None else "0"
        return None

    # ---- body ------------------------------------------------------------
    def read(self):
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        cursor = QTextCursor(self.qdoc)
        self._items(cursor, self.word.element.body, self.word.part, fresh=True,
                    Paragraph=Paragraph, Table=Table)

    def _items(self, cursor, parent, part, fresh, Paragraph, Table):
        """Write the paragraphs and tables under `parent` at `cursor`.
        fresh: the cursor's block is still empty and unused. Returns fresh."""
        for child in parent.iterchildren():
            if child.tag == _qn("w:p"):
                self._paragraph(cursor, Paragraph(child, _PartHolder(part)), part, fresh)
                fresh = False
            elif child.tag == _qn("w:tbl"):
                cursor.movePosition(QTextCursor.EndOfBlock)  # the table goes after this block
                self._table(cursor, child, part, Paragraph, Table)
                fresh = True
            elif child.tag == _qn("w:sdt"):  # content controls (e.g. a table of contents)
                content = child.find(_qn("w:sdtContent"))
                if content is not None:
                    fresh = self._items(cursor, content, part, fresh, Paragraph, Table)
        return fresh

    def _paragraph(self, cursor, para, part, fresh):
        style = para.style
        chain = self._chain(style)
        name = style.name if style is not None else "Normal"
        bf = QTextBlockFormat()
        bf.setProperty(STYLE_PROP, name)
        level = None
        if name == "Title":
            level = 0
        elif name.startswith("Heading ") and name[8:].strip().isdigit():
            level = int(name[8:])
        if level is not None:
            bf.setHeadingLevel(max(1, level))
        self._block_layout(bf, para, chain)
        if self.page_break or self._para_attr(para, "page_break_before", chain):
            bf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
            self.page_break = False
        base = self._char_format(None, chain, level)
        if fresh:
            cursor.setBlockFormat(bf)
            cursor.setBlockCharFormat(base)
        else:
            cursor.insertBlock(bf, base)
        num = self._num_pr(para, chain)
        if num is not None:
            bf.setProperty(NUM_PROP, f"{num[0]}:{num[1]}")
            cursor.setBlockFormat(bf)
            self._join_list(cursor, num)
        for run, href in self._runs(para._p, para, part):
            rpr = run._r.rPr
            run_chain = self._chain(run.style) if rpr is not None and rpr.rStyle is not None else []
            self._run(cursor, run, run_chain + chain, part, level, href)

    def _block_layout(self, bf, para, chain):
        from docx.enum.text import WD_ALIGN_PARAGRAPH as A

        align = self._para_attr(para, "alignment", chain)
        bf.setAlignment({A.CENTER: Qt.AlignHCenter, A.RIGHT: Qt.AlignRight,
                         A.JUSTIFY: Qt.AlignJustify}.get(align, Qt.AlignLeft))
        for attr, setter in (("left_indent", bf.setLeftMargin), ("right_indent", bf.setRightMargin),
                             ("first_line_indent", bf.setTextIndent), ("space_before", bf.setTopMargin),
                             ("space_after", bf.setBottomMargin)):
            value = self._para_attr(para, attr, chain)
            if value is not None:
                setter(value.pt * PX_PER_PT)
        spacing = self._para_attr(para, "line_spacing", chain)
        if isinstance(spacing, float) and 0.5 <= spacing <= 4:
            bf.setLineHeight(spacing * 100, _PROPORTIONAL)
        ppr = para._p.pPr
        if ppr is not None:
            shd = ppr.find(_qn("w:shd"))
            fill = shd.get(_qn("w:fill")) if shd is not None else None
            if fill and fill.lower() not in ("auto", "ffffff"):
                bf.setBackground(QColor(f"#{fill}"))
            bdr = ppr.find(_qn("w:pBdr"))
            if bdr is not None and bdr.find(_qn("w:bottom")) is not None and not para.text.strip():
                bf.setProperty(HR_PROP, QTextLength(QTextLength.PercentageLength, 100))

    def _join_list(self, cursor, num):
        existing = self.lists.get(num)
        if existing is not None:
            existing.add(cursor.block())
        else:
            fmt = QTextListFormat()
            level = int(num[1]) if num[1].isdigit() else 0
            kind, text, start = self._numbering.get(num, ("bullet", "", 1))
            if kind == "bullet":
                fmt.setStyle({"o": QTextListFormat.ListCircle, "\u25e6": QTextListFormat.ListCircle,
                              "\u00a7": QTextListFormat.ListSquare, "\u25aa": QTextListFormat.ListSquare,
                              "\u25a0": QTextListFormat.ListSquare, "\uf0a7": QTextListFormat.ListSquare,
                              "\u2022": QTextListFormat.ListDisc, "\uf0b7": QTextListFormat.ListDisc}.get(
                    text, _BULLETS[level % 3]))
            else:
                fmt.setStyle(_NUMBER_STYLES.get(kind, QTextListFormat.ListDecimal))
                if text.count("%") == 1:   # "(%1)" -> "(" and ")"; "%1." -> "."
                    before, after = text.split("%")
                    fmt.setNumberPrefix(before)
                    fmt.setNumberSuffix(after[1:] if after[:1].isdigit() else after)
                if start != 1:
                    fmt.setStart(start)
            fmt.setIndent(level + 1)
            self.lists[num] = cursor.createList(fmt)
        # the list takes over the indent Word gave the paragraph
        bf = cursor.blockFormat()
        bf.setLeftMargin(0)
        bf.setTextIndent(0)
        cursor.setBlockFormat(bf)

    def _runs(self, element, para, part, href=None):
        from docx.text.run import Run

        for child in element.iterchildren():
            if child.tag == _qn("w:r"):
                yield Run(child, para), href
            elif child.tag == _qn("w:hyperlink"):
                rid = child.get(_qn("r:id"))
                link = None
                if rid and rid in part.rels:
                    link = part.rels[rid].target_ref
                elif child.get(_qn("w:anchor")):
                    link = "#" + child.get(_qn("w:anchor"))
                yield from self._runs(child, para, part, link)
            elif child.tag in (_qn("w:ins"), _qn("w:smartTag"), _qn("w:fldSimple"), _qn("w:customXml")):
                yield from self._runs(child, para, part, href)
            elif child.tag == _qn("w:sdt"):
                content = child.find(_qn("w:sdtContent"))
                if content is not None:
                    yield from self._runs(content, para, part, href)

    def _char_format(self, run, chain, level):
        """The look of a run (or, with run=None, of the paragraph itself)
        resolved through `chain`: the run's and the paragraph's styles."""
        cf = QTextCharFormat()
        family, size = self._defaults

        def attr(name):
            value = getattr(run.font, name) if run is not None else None
            return value if value is not None else self._first(chain, name)

        family = attr("name") or family
        sz = attr("size")
        size = sz.pt if sz is not None else size
        bold, italic = attr("bold"), attr("italic")
        color = None
        if run is not None:
            try:
                if run.font.color.type is not None and run.font.color.rgb:
                    color = run.font.color.rgb
            except (AttributeError, ValueError):
                color = None
            if attr("underline"):
                cf.setFontUnderline(True)
                rpr = run._r.rPr
                u = rpr.find(_qn("w:u")) if rpr is not None else None
                kind = u.get(_qn("w:val")) if u is not None else "single"
                if _UNDERLINES.get(kind, QTextCharFormat.SingleUnderline) != QTextCharFormat.SingleUnderline:
                    cf.setUnderlineStyle(_UNDERLINES[kind])
            if attr("all_caps"):
                cf.setFontCapitalization(QFont.AllUppercase)
            elif attr("small_caps"):
                cf.setFontCapitalization(QFont.SmallCaps)
            if attr("strike"):
                cf.setFontStrikeOut(True)
            if attr("superscript"):
                cf.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
            elif attr("subscript"):
                cf.setVerticalAlignment(QTextCharFormat.AlignSubScript)
            self._highlight(cf, run)
        if color is None:
            color = self._first_color(chain)
        if level is not None and bold is None:
            bold = True
        if level is not None and size == self._defaults[1]:
            size = _HEADING_SIZES.get(level, size)
        cf.setFontFamilies([family])
        cf.setFontPointSize(size)
        if bold:
            cf.setFontWeight(QFont.Bold)
        if italic:
            cf.setFontItalic(True)
        if color is not None:
            cf.setForeground(QColor(f"#{color}"))
        return cf

    @staticmethod
    def _highlight(cf, run):
        rpr = run._r.rPr
        if rpr is None:
            return
        hl = rpr.find(_qn("w:highlight"))
        if hl is not None and hl.get(_qn("w:val")) in _HIGHLIGHTS:
            cf.setBackground(QColor(_HIGHLIGHTS[hl.get(_qn("w:val"))]))
            return
        shd = rpr.find(_qn("w:shd"))
        fill = shd.get(_qn("w:fill")) if shd is not None else None
        if fill and fill.lower() not in ("auto", "ffffff"):
            cf.setBackground(QColor(f"#{fill}"))

    def _run(self, cursor, run, chain, part, level, href):
        cf = self._char_format(run, chain, level)
        if href:
            cf.setAnchor(True)
            cf.setAnchorHref(href)
            cf.setFontUnderline(True)
            if not cf.hasProperty(QTextFormat.ForegroundBrush):
                cf.setForeground(QColor("#0563c1"))
        text = []

        def flush():
            if text:
                cursor.insertText("".join(text), cf)
                text.clear()

        for child in run._r.iterchildren():
            tag = child.tag
            if tag == _qn("w:t"):
                text.append(child.text or "")
            elif tag == _qn("w:tab"):
                text.append("\t")
            elif tag in (_qn("w:br"), _qn("w:cr")):
                if child.get(_qn("w:type")) == "page":
                    self.page_break = True
                else:
                    text.append(" ")  # line break inside the paragraph
            elif tag == _qn("w:noBreakHyphen"):
                text.append("‑")
            elif tag == _qn("w:sym"):
                try:
                    text.append(chr(int(child.get(_qn("w:char")), 16)))
                except (TypeError, ValueError):
                    pass
            elif tag in (_qn("w:drawing"), _qn("w:pict")):
                flush()
                self._image(cursor, child, part)
        flush()

    def _image(self, cursor, element, part):
        blip = next((e for e in element.iter() if e.tag in (_qn("a:blip"), _VML_IMAGEDATA)), None)
        if blip is None:
            return
        rid = blip.get(_qn("r:embed")) or blip.get(_qn("r:id"))
        try:
            blob = part.related_parts[rid].blob
        except (KeyError, AttributeError):
            return
        image = QImage.fromData(blob)
        if image.isNull():
            return
        self.images += 1
        url = f"docx-image-{self.images}"
        self.qdoc.addResource(QTextDocument.ImageResource, QUrl(url), image)
        fmt = QTextImageFormat()
        fmt.setName(url)
        extent = next((e for e in element.iter() if e.tag == _qn("wp:extent")), None)
        if extent is not None:
            fmt.setWidth(int(extent.get("cx")) / EMU_PER_PX)
            fmt.setHeight(int(extent.get("cy")) / EMU_PER_PX)
        else:
            fmt.setWidth(image.width())
            fmt.setHeight(image.height())
        cursor.insertImage(fmt)

    def _table(self, cursor, tbl, part, Paragraph, Table):
        rows = tbl.findall(_qn("w:tr"))
        grid = tbl.find(_qn("w:tblGrid"))
        widths = [int(g.get(_qn("w:w")) or 0) for g in grid.findall(_qn("w:gridCol"))] if grid is not None else []
        cells, n_cols = [], len(widths)
        for r, tr in enumerate(rows):
            col = 0
            for tc in tr.findall(_qn("w:tc")):
                tcpr = tc.find(_qn("w:tcPr"))
                span, vmerge, fill = 1, None, None
                if tcpr is not None:
                    gs = tcpr.find(_qn("w:gridSpan"))
                    if gs is not None:
                        span = max(1, int(gs.get(_qn("w:val")) or 1))
                    vm = tcpr.find(_qn("w:vMerge"))
                    if vm is not None:
                        vmerge = vm.get(_qn("w:val")) or "continue"
                    shd = tcpr.find(_qn("w:shd"))
                    if shd is not None and (shd.get(_qn("w:fill")) or "auto").lower() not in ("auto", "ffffff"):
                        fill = shd.get(_qn("w:fill"))
                cells.append((r, col, span, tc, vmerge, fill))
                col += span
            n_cols = max(n_cols, col)
        if not rows or not n_cols:
            return
        fmt = QTextTableFormat()
        fmt.setBorder(1)
        fmt.setBorderBrush(QBrush(QColor("#9aa0a6")))
        fmt.setBorderCollapse(True)
        fmt.setCellPadding(4)
        fmt.setCellSpacing(0)
        fmt.setWidth(QTextLength(QTextLength.PercentageLength, 100))
        if len(widths) == n_cols and sum(widths):
            fmt.setColumnWidthConstraints([QTextLength(QTextLength.PercentageLength, 100 * w / sum(widths))
                                           for w in widths])
        table = cursor.insertTable(len(rows), n_cols, fmt)
        for r, col, span, tc, vmerge, fill in cells:
            if vmerge == "continue" or col >= n_cols:
                continue
            cell = table.cellAt(r, col)
            if fill:
                cell_fmt = cell.format().toTableCellFormat()
                cell_fmt.setBackground(QColor(f"#{fill}"))
                cell.setFormat(cell_fmt)
            self._items(cell.firstCursorPosition(), tc, part, fresh=True, Paragraph=Paragraph, Table=Table)
        for r, col, span, tc, vmerge, fill in cells:
            if vmerge == "continue" or col >= n_cols:
                continue
            row_span = 1
            if vmerge == "restart":
                for below in range(r + 1, len(rows)):
                    if any(c[0] == below and c[1] == col and c[4] == "continue" for c in cells):
                        row_span += 1
                    else:
                        break
            if row_span > 1 or span > 1:
                table.mergeCells(r, col, row_span, min(span, n_cols - col))
        after = table.lastCursorPosition()
        after.movePosition(QTextCursor.NextBlock)
        cursor.setPosition(after.position())


class _PartHolder:
    """python-docx Paragraph/Run look their part up through their parent."""

    def __init__(self, part):
        self.part = part


def load_docx(path):
    """(QTextDocument, PageSetup) for a .docx file."""
    from docx import Document

    word = Document(path)
    qdoc = QTextDocument()
    page = PageSetup()
    try:
        section = word.sections[0]
        page = PageSetup(section.page_width.pt, section.page_height.pt,
                         (section.left_margin.pt, section.top_margin.pt, section.right_margin.pt,
                          section.bottom_margin.pt))
    except (IndexError, AttributeError, TypeError):
        pass
    reader = _Reader(word, qdoc)
    family, size = reader._defaults
    qdoc.setDefaultFont(QFont(family, round(size)))
    reader.read()
    apply_page(qdoc, page)
    qdoc.setModified(False)
    return qdoc, page


def apply_page(qdoc, page):
    fmt = qdoc.rootFrame().frameFormat()
    left, top, right, bottom = (m * PX_PER_PT for m in page.margins_pt)
    fmt.setLeftMargin(left)
    fmt.setTopMargin(top)
    fmt.setRightMargin(right)
    fmt.setBottomMargin(bottom)
    qdoc.rootFrame().setFrameFormat(fmt)
    qdoc.setDocumentMargin(0)


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------

class _Writer:
    def __init__(self, qdoc, word):
        self.qdoc = qdoc
        self.word = word
        self.styles = {s.name for s in word.styles}
        self._num_ids = {}   # a QTextList's objectIndex -> the Word numId written for it

    def write(self):
        body = self.word.element.body
        for child in list(body):
            if child.tag != _qn("w:sectPr"):
                body.remove(child)
        self._frame(self.qdoc.rootFrame(), self.word, first_paragraph=None)

    def _frame(self, frame, container, first_paragraph):
        """Write a frame's blocks and tables into a Document or table cell.
        first_paragraph: an empty paragraph to fill before adding new ones."""
        it = frame.begin()
        while not it.atEnd():
            child = it.currentFrame()
            if isinstance(child, QTextTable):
                self._table(child, container)
                first_paragraph = None
            elif child is not None:  # some other frame: write its contents inline
                self._frame(child, container, first_paragraph)
                first_paragraph = None
            else:
                block = it.currentBlock()
                if block.isValid():
                    para = first_paragraph if first_paragraph is not None else container.add_paragraph()
                    first_paragraph = None
                    self._paragraph(block, para)
            it += 1

    def _style_for(self, block):
        """The Word style to give a block: the one it was read with, else one
        that matches what it is now (a heading level, a new list)."""
        bf = block.blockFormat()
        name = bf.property(STYLE_PROP)
        name = name if isinstance(name, str) and name in self.styles else None
        is_heading_style = name is not None and (name == "Title" or name.startswith("Heading"))
        level = bf.headingLevel()
        if level:
            if is_heading_style:
                return name
            wanted = f"Heading {min(level, 9)}"
            return wanted if wanted in self.styles else None
        if block.textList() is not None and not bf.property(NUM_PROP):
            return "List Paragraph" if "List Paragraph" in self.styles else name
        return None if is_heading_style else name

    def _paragraph(self, block, para):
        from docx.enum.text import WD_ALIGN_PARAGRAPH as A
        from docx.oxml import OxmlElement
        from docx.shared import Pt

        bf = block.blockFormat()
        style = self._style_for(block)
        if style:
            try:
                para.style = self.word.styles[style]
            except KeyError:
                pass
        pf = para.paragraph_format
        align = bf.alignment() & Qt.AlignHorizontal_Mask
        pf.alignment = {Qt.AlignHCenter: A.CENTER, Qt.AlignRight: A.RIGHT, Qt.AlignJustify: A.JUSTIFY}.get(
            align, A.LEFT)
        for value, attr in ((bf.leftMargin(), "left_indent"), (bf.rightMargin(), "right_indent"),
                            (bf.textIndent(), "first_line_indent"), (bf.topMargin(), "space_before"),
                            (bf.bottomMargin(), "space_after")):
            if value:
                setattr(pf, attr, Pt(value / PX_PER_PT))
        if bf.lineHeightType() == _PROPORTIONAL and bf.lineHeight() not in (0, 100):
            pf.line_spacing = bf.lineHeight() / 100
        if bf.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
            pf.page_break_before = True
        num = bf.property(NUM_PROP)
        if block.textList() is not None and isinstance(num, str) and ":" in num:
            num_id, ilvl = num.split(":", 1)
            ppr = para._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            lvl_el = OxmlElement("w:ilvl")
            lvl_el.set(_qn("w:val"), ilvl)
            id_el = OxmlElement("w:numId")
            id_el.set(_qn("w:val"), num_id)
            num_pr.append(lvl_el)
            num_pr.append(id_el)
            ppr.append(num_pr)
        elif block.textList() is not None:   # a list made here: give it Word numbering of its own
            lst = block.textList()
            num_id = self._numbering_for(lst)
            if num_id is None:
                para.add_run("\u2022 " if lst.format().style() in _BULLETS else lst.itemText(block) + " ")
            else:
                ppr = para._p.get_or_add_pPr()
                num_pr = OxmlElement("w:numPr")
                lvl_el = OxmlElement("w:ilvl")
                lvl_el.set(_qn("w:val"), str(max(0, lst.format().indent() - 1)))
                id_el = OxmlElement("w:numId")
                id_el.set(_qn("w:val"), str(num_id))
                num_pr.append(lvl_el)
                num_pr.append(id_el)
                ppr.append(num_pr)
        if bf.background().style() != Qt.NoBrush:
            color = bf.background().color()
            shd = OxmlElement("w:shd")
            shd.set(_qn("w:val"), "clear")
            shd.set(_qn("w:color"), "auto")
            shd.set(_qn("w:fill"), f"{color.red():02X}{color.green():02X}{color.blue():02X}")
            para._p.get_or_add_pPr().append(shd)
        if bf.hasProperty(HR_PROP):
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "auto")):
                bottom.set(_qn(key), value)
            bdr.append(bottom)
            para._p.get_or_add_pPr().append(bdr)
        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid():
                self._fragment(frag, para)
            it += 1

    def _numbering_for(self, lst):
        """The numId of a Word numbering definition that looks like `lst`
        (one per list, so each list counts from its own start)."""
        from docx.oxml import OxmlElement

        key = lst.objectIndex()
        if key in self._num_ids:
            return self._num_ids[key]
        try:
            root = self.word.part.numbering_part.element
        except (NotImplementedError, KeyError, AttributeError):
            return None
        fmt = lst.format()
        level = max(0, fmt.indent() - 1)
        ids = [int(e.get(_qn("w:abstractNumId"))) for e in root.findall(_qn("w:abstractNum"))
               if (e.get(_qn("w:abstractNumId")) or "").isdigit()]
        abstract_id = max(ids, default=-1) + 1
        num_ids = [int(e.get(_qn("w:numId"))) for e in root.findall(_qn("w:num"))
                   if (e.get(_qn("w:numId")) or "").isdigit()]
        num_id = max(num_ids, default=0) + 1

        def el(tag, **attrs):
            e = OxmlElement(tag)
            for k, v in attrs.items():
                e.set(_qn(f"w:{k}"), str(v))
            return e

        an = el("w:abstractNum", abstractNumId=abstract_id)
        an.append(el("w:multiLevelType", val="hybridMultilevel"))
        for ilvl in range(9):
            lvl = el("w:lvl", ilvl=ilvl)
            style = fmt.style() if ilvl == level else (_BULLETS[ilvl % 3] if fmt.style() in _BULLETS
                                                         else QTextListFormat.ListDecimal)
            lvl.append(el("w:start", val=fmt.start() if ilvl == level and fmt.start() > 0 else 1))
            font = None
            if style in _BULLETS:
                char, font = _BULLET_CHARS[style]
                lvl.append(el("w:numFmt", val="bullet"))
                lvl.append(el("w:lvlText", val=char))
            else:
                lvl.append(el("w:numFmt", val=_NUM_FORMATS.get(style, "decimal")))
                prefix, suffix = (fmt.numberPrefix(), fmt.numberSuffix()) if ilvl == level else ("", ".")
                lvl.append(el("w:lvlText", val=f"{prefix}%{ilvl + 1}{suffix}"))
            lvl.append(el("w:lvlJc", val="left"))
            ppr = el("w:pPr")
            ppr.append(el("w:ind", left=720 * (ilvl + 1), hanging=360))
            lvl.append(ppr)
            if font:
                rpr = el("w:rPr")
                rpr.append(el("w:rFonts", ascii=font, hAnsi=font, hint="default"))
                lvl.append(rpr)
            an.append(lvl)
        first_num = root.find(_qn("w:num"))
        if first_num is not None:
            first_num.addprevious(an)      # abstract definitions come before the numbers
        else:
            root.append(an)
        num = el("w:num", numId=num_id)
        num.append(el("w:abstractNumId", val=abstract_id))
        root.append(num)
        self._num_ids[key] = num_id
        return num_id

    def _fragment(self, frag, para):
        cf = frag.charFormat()
        if cf.isImageFormat():
            self._picture(cf.toImageFormat(), para, frag.length())
            return
        text = frag.text().replace(" ", "\n").replace(" ", "\n").replace("￼", "")
        if not text:
            return
        run = para.add_run(text)
        self._run_format(run, cf)
        if cf.isAnchor() and cf.anchorHref() and not cf.anchorHref().startswith("#"):
            self._hyperlink(para, run, cf.anchorHref())

    def _run_format(self, run, cf):
        from docx.oxml import OxmlElement
        from docx.shared import Pt, RGBColor

        font = run.font
        if cf.fontWeight() >= QFont.DemiBold:
            font.bold = True
        if cf.fontItalic():
            font.italic = True
        if cf.fontUnderline():
            from docx.enum.text import WD_UNDERLINE

            font.underline = {"dotted": WD_UNDERLINE.DOTTED, "dash": WD_UNDERLINE.DASH,
                              "dotDash": WD_UNDERLINE.DOT_DASH, "dotDotDash": WD_UNDERLINE.DOT_DOT_DASH,
                              "wave": WD_UNDERLINE.WAVY}.get(_UNDERLINE_OUT.get(cf.underlineStyle()), True)
        if cf.fontCapitalization() == QFont.AllUppercase:
            font.all_caps = True
        elif cf.fontCapitalization() == QFont.SmallCaps:
            font.small_caps = True
        if cf.fontStrikeOut():
            font.strike = True
        valign = cf.verticalAlignment()
        if valign == QTextCharFormat.AlignSuperScript:
            font.superscript = True
        elif valign == QTextCharFormat.AlignSubScript:
            font.subscript = True
        families = cf.fontFamilies()
        if families:
            font.name = families[0]
        if cf.fontPointSize() > 0:
            font.size = Pt(round(cf.fontPointSize() * 2) / 2)
        if cf.hasProperty(QTextFormat.ForegroundBrush):
            color = cf.foreground().color()
            font.color.rgb = RGBColor(color.red(), color.green(), color.blue())
        if cf.hasProperty(QTextFormat.BackgroundBrush) and cf.background().style() != Qt.NoBrush:
            color = cf.background().color()
            name = next((k for k, v in _HIGHLIGHTS.items() if v == color.name()), None)
            if name:   # one of Word's highlighter colours: a real highlight
                hl = OxmlElement("w:highlight")
                hl.set(_qn("w:val"), name)
                run._r.get_or_add_rPr().append(hl)
                return
            shd = OxmlElement("w:shd")
            shd.set(_qn("w:val"), "clear")
            shd.set(_qn("w:color"), "auto")
            shd.set(_qn("w:fill"), f"{color.red():02X}{color.green():02X}{color.blue():02X}")
            run._r.get_or_add_rPr().append(shd)

    @staticmethod
    def _hyperlink(para, run, href):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement

        rid = para.part.relate_to(href, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(_qn("r:id"), rid)
        run._r.addprevious(link)
        link.append(run._r)

    def _picture(self, fmt, para, count):
        from docx.shared import Pt

        resource = self.qdoc.resource(QTextDocument.ImageResource, QUrl(fmt.name()))
        if isinstance(resource, QPixmap):
            resource = resource.toImage()
        image = resource if isinstance(resource, QImage) else QImage()
        if image.isNull():
            image = QImage(fmt.name())  # a file path
        if image.isNull():
            return
        data = QByteArray()
        buf = QBuffer(data)
        buf.open(QIODevice.WriteOnly)
        image.save(buf, "PNG")
        buf.close()
        width = fmt.width() or image.width()
        for _ in range(max(1, count)):
            para.add_run().add_picture(io.BytesIO(bytes(data)), width=Pt(width / PX_PER_PT))

    def _table(self, table, container):
        rows, cols = table.rows(), table.columns()
        word_table = container.add_table(rows=rows, cols=cols)
        if "Table Grid" in self.styles:
            word_table.style = self.word.styles["Table Grid"]
        for r in range(rows):
            for c in range(cols):
                cell = table.cellAt(r, c)
                if cell.row() != r or cell.column() != c:
                    continue  # covered by a merged cell
                target = word_table.cell(r, c)
                if cell.rowSpan() > 1 or cell.columnSpan() > 1:
                    target = target.merge(word_table.cell(r + cell.rowSpan() - 1, c + cell.columnSpan() - 1))
                bg = cell.format().background()
                if bg.style() != Qt.NoBrush:
                    from docx.oxml import OxmlElement

                    shd = OxmlElement("w:shd")
                    shd.set(_qn("w:val"), "clear")
                    shd.set(_qn("w:fill"), bg.color().name()[1:].upper())
                    target._tc.get_or_add_tcPr().append(shd)
                first = target.paragraphs[0]
                for extra in target.paragraphs[1:]:  # merge() joins the cells' paragraphs
                    extra._p.getparent().remove(extra._p)
                self._cell(cell, target, first)

    def _cell(self, cell, target, first):
        it = cell.begin()
        while not it.atEnd():
            child = it.currentFrame()
            if isinstance(child, QTextTable):
                self._table(child, target)
                first = None
            else:
                block = it.currentBlock()
                if block.isValid():
                    para = first if first is not None else target.add_paragraph()
                    first = None
                    self._paragraph(block, para)
            it += 1


def save_docx(qdoc, path, base=None, page=None):
    """Write `qdoc` to `path`. base: the .docx it came from, whose styles,
    headers, footers and page setup are kept."""
    from docx import Document
    from docx.shared import Pt

    from docx.enum.section import WD_ORIENT

    word = Document(base) if base and os.path.isfile(base) else Document()
    if page is not None:   # the size, orientation and margins set here (or read from the file)
        for section in word.sections:
            section.orientation = WD_ORIENT.LANDSCAPE if page.width_pt > page.height_pt else WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Pt(page.width_pt), Pt(page.height_pt)
            (section.left_margin, section.top_margin, section.right_margin,
             section.bottom_margin) = (Pt(m) for m in page.margins_pt)
    _Writer(qdoc, word).write()
    folder = os.path.dirname(os.path.abspath(path))
    tmp = os.path.join(folder, f".~{os.path.basename(path)}.tmp")
    try:
        word.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
