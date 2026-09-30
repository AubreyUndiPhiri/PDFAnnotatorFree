"""Convert a PDF to an editable Word document (.docx).

- Typed pages (with a text layer) go through pdf2docx, which rebuilds the
  paragraphs, fonts, colours, tables and images.
- Scanned pages (a picture of the page, no text) are read with OCR and
  written as editable headings, paragraphs and tables: Tesseract when it is
  installed (through PyMuPDF), otherwise the GOT-OCR 2.0 model, or TrOCR
  for handwriting (ocr_models.py).
- Any page that still can't be converted is kept as a picture of the page,
  so nothing is ever left out of the Word file.

Annotations (highlights, ink, stamps, text boxes...) are burnt into a copy
of the pages first, so the Word file shows what the PDF shows.

No Qt here: the UI runs PdfToDocx.run() in a worker thread.
"""
import io
import logging
import os
import re
import statistics
import tempfile
from dataclasses import dataclass, field

import pymupdf as fitz

OCR_AUTO = "auto"
OCR_TESSERACT = "tesseract"
OCR_GOT = "got"
OCR_TROCR = "trocr"
OCR_NONE = "none"
OCR_CHOICES = {
    OCR_AUTO: "Automatic (best OCR available)",
    OCR_TESSERACT: "Tesseract OCR",
    OCR_GOT: "GOT-OCR 2.0 model (printed text, tables)",
    OCR_TROCR: "TrOCR model (handwriting)",
    OCR_NONE: "No OCR (keep scanned pages as pictures)",
}
ENGINE_LABELS = {OCR_TESSERACT: "Tesseract", OCR_GOT: "GOT-OCR 2.0", OCR_TROCR: "TrOCR"}

OCR_DPI = 300
PICTURE_DPI = 150
OCR_MARGIN_PT = 54  # 0.75 in, for pages rebuilt from OCR text
PICTURE_MARGIN_PT = 18


class ConversionCancelled(Exception):
    pass


class ConversionError(Exception):
    pass


@dataclass
class ConversionReport:
    pages: int = 0
    converted: list = field(default_factory=list)  # 1-based page numbers, by kind
    ocr: list = field(default_factory=list)
    pictures: list = field(default_factory=list)
    ocr_engine: str | None = None
    warnings: list = field(default_factory=list)

    def summary(self) -> str:
        parts = []
        if self.converted:
            parts.append(f"{len(self.converted)} page(s) converted to editable text")
        if self.ocr:
            parts.append(f"{len(self.ocr)} scanned page(s) read with {ENGINE_LABELS.get(self.ocr_engine, 'OCR')}")
        if self.pictures:
            parts.append(f"{len(self.pictures)} page(s) kept as pictures ({page_list(self.pictures)})")
        return "; ".join(parts) + "." if parts else "No pages were converted."


def page_list(numbers) -> str:
    """[1, 2, 3, 7] -> '1-3, 7'"""
    out, run = [], []
    for n in sorted(numbers):
        if run and n == run[-1] + 1:
            run.append(n)
            continue
        if run:
            out.append(f"{run[0]}-{run[-1]}" if len(run) > 1 else str(run[0]))
        run = [n]
    if run:
        out.append(f"{run[0]}-{run[-1]}" if len(run) > 1 else str(run[0]))
    return ", ".join(out)


def parse_page_range(text: str, page_count: int) -> list[int]:
    """'1-3, 5, 8-' -> 0-based indexes. Empty text means every page."""
    text = (text or "").strip().lower()
    if not text or text == "all":
        return list(range(page_count))
    pages = []
    for part in re.split(r"[,;\s]+", text):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)\s*-\s*(\d*)|(\d+)", part)
        if not m:
            raise ValueError(f'"{part}" is not a page number or range (for example 1-3, 5).')
        if m.group(3):
            first = last = int(m.group(3))
        else:
            first = int(m.group(1) or 1)
            last = int(m.group(2) or page_count)
        if first < 1 or last > page_count or first > last:
            raise ValueError(f'Pages "{part}" are outside 1-{page_count}.')
        pages.extend(i for i in range(first - 1, last) if i not in pages)
    return pages


def is_scanned(page) -> bool:
    """A page with no text at all whose images cover a good part of it."""
    if page.get_text("text").strip():
        return False
    area = abs(page.rect) or 1
    covered = sum(abs(fitz.Rect(info["bbox"]) & page.rect) for info in page.get_image_info())
    return covered >= 0.3 * area


def tesseract_available() -> bool:
    try:
        return bool(fitz.get_tessdata())
    except Exception:  # noqa: BLE001 - raises when Tesseract isn't installed
        return False


# --------------------------------------------------------------------------
# OCR output -> blocks: ("heading", level, text) / ("para", text) / ("table", rows)
# --------------------------------------------------------------------------

def tesseract_blocks(page) -> list:
    tp = page.get_textpage_ocr(dpi=OCR_DPI, full=True)
    data = page.get_text("dict", textpage=tp, sort=True)
    paragraphs = []
    for block in data["blocks"]:
        if block.get("type") != 0:
            continue
        lines, sizes = [], []
        for line in block["lines"]:
            text = " ".join(s["text"].strip() for s in line["spans"] if s["text"].strip())
            if text:
                lines.append(text)
                sizes.extend(s["size"] for s in line["spans"] if s["text"].strip())
        if lines:
            paragraphs.append((_join_lines(lines), max(sizes)))
    if not paragraphs:
        return []
    body = statistics.median(size for _, size in paragraphs)
    blocks = []
    for text, size in paragraphs:
        if size >= body * 1.35 and len(text) < 120:
            blocks.append(("heading", 1 if size >= body * 1.8 else 2, text))
        else:
            blocks.append(("para", text))
    return blocks


def _join_lines(lines):
    out = lines[0]
    for line in lines[1:]:
        out = out[:-1] + line if out.endswith("-") and out[-2:-1].isalpha() else f"{out} {line}"
    return out


_TABULAR = re.compile(r"\\begin\{tabular\}(?:\{[^}]*\})?(.*?)\\end\{tabular\}", re.S)
_HEADING = re.compile(r"\\(title|section|subsection|subsubsection)\*?\{(.*?)\}", re.S)
_TEXT_CMD = re.compile(r"\\(?:textbf|textit|emph|underline|text|mathrm|caption|footnote|author|date|"
                       r"multirow\{[^}]*\}\{[^}]*\}|multicolumn\{[^}]*\}\{[^}]*\})\{([^{}]*)\}")
_DROP_CMD = re.compile(r"\\(?:hline|centering|cline\{[^}]*\}|begin\{(?:table|center|figure)\}(?:\[[^]]*\])?"
                       r"|end\{(?:table|center|figure)\}|maketitle|noindent|newline)")
_UNESCAPE = {r"\%": "%", r"\&": "&", r"\_": "_", r"\$": "$", r"\#": "#", r"\{": "{", r"\}": "}", "~": " "}


def clean_latex(text: str) -> str:
    """The light LaTeX that GOT-OCR writes -> plain text (formulas keep their
    LaTeX source, without the \\( \\) delimiters)."""
    text = _DROP_CMD.sub("", text)
    for _ in range(4):  # nested \textbf{\textit{...}}
        text, n = _TEXT_CMD.subn(r"\1", text)
        if not n:
            break
    text = re.sub(r"\\[()\[\]]", "", text)
    for escaped, plain in _UNESCAPE.items():
        text = text.replace(escaped, plain)
    return re.sub(r"[ \t]+", " ", text).strip()


def got_blocks(text: str) -> list:
    blocks, pos = [], 0
    for m in _TABULAR.finditer(text):
        _got_text_blocks(text[pos:m.start()], blocks)
        rows = []
        for raw_row in re.split(r"\\\\", m.group(1)):
            cells = [clean_latex(c) for c in re.split(r"(?<!\\)&", raw_row)]
            if any(cells):
                rows.append(cells)
        if rows:
            blocks.append(("table", rows))
        pos = m.end()
    _got_text_blocks(text[pos:], blocks)
    return blocks


def _got_text_blocks(text, blocks):
    pos = 0
    for m in _HEADING.finditer(text):
        _plain_blocks(text[pos:m.start()], blocks)
        heading = clean_latex(m.group(2))
        if heading:
            level = {"title": 0, "section": 1, "subsection": 2}.get(m.group(1), 3)
            blocks.append(("heading", level, heading))
        pos = m.end()
    _plain_blocks(text[pos:], blocks)


def _plain_blocks(text, blocks):
    for line in text.splitlines():
        line = clean_latex(line)
        if line:
            blocks.append(("para", line))


def trocr_blocks(text: str) -> list:
    return [("para", line.strip()) for line in text.splitlines() if line.strip()]


# --------------------------------------------------------------------------
# Conversion
# --------------------------------------------------------------------------

class PdfToDocx:
    """Convert a PDF (a path, bytes or an open fitz.Document) to .docx.

        PdfToDocx(src, pages=None, ocr=OCR_AUTO).run("out.docx")

    progress(fraction, text) is called as it goes; cancel() can be called
    from another thread and makes run() raise ConversionCancelled."""

    def __init__(self, source, pages=None, password=None, ocr=OCR_AUTO, include_annotations=True,
                 keep_scan_images=False, progress=None):
        self.source = source
        self.pages = pages
        self.password = password
        self.ocr = ocr
        self.include_annotations = include_annotations
        self.keep_scan_images = keep_scan_images
        self.progress = progress or (lambda fraction, text: None)
        self.report = ConversionReport()
        self._cancelled = False
        self._model = None

    def cancel(self):
        self._cancelled = True
        if self._model is not None:
            self._model.kill()

    def _check_cancel(self):
        if self._cancelled:
            raise ConversionCancelled()

    # ------------------------------------------------------------ source
    def _open_source(self) -> fitz.Document:
        if isinstance(self.source, fitz.Document):
            data = self.source.tobytes()
        elif isinstance(self.source, (bytes, bytearray)):
            data = bytes(self.source)
        else:
            with open(self.source, "rb") as f:
                data = f.read()
        try:
            doc = fitz.open(stream=data, filetype="pdf")
        except Exception as e:  # noqa: BLE001 - PyMuPDF raises several types for bad files
            raise ConversionError(f"This file could not be read as a PDF: {e}") from e
        if doc.needs_pass and not doc.authenticate(self.password or ""):
            raise ConversionError("This PDF is password protected. Open it with its password first.")
        if doc.page_count == 0:
            raise ConversionError("This PDF has no pages.")
        # Work on a plain, decrypted copy with the annotations burnt in
        doc = fitz.open(stream=doc.tobytes(garbage=1), filetype="pdf")
        if self.include_annotations:
            try:
                doc.bake(annots=True, widgets=True)
            except Exception as e:  # noqa: BLE001 - keep going without the markup
                self.report.warnings.append(f"Annotations could not be included: {e}")
        return doc

    # ------------------------------------------------------------ run
    def run(self, out_path):
        from docx import Document

        self.progress(0.0, "Opening the PDF...")
        doc = self._open_source()
        indexes = list(range(doc.page_count)) if self.pages is None else [i for i in self.pages
                                                                          if 0 <= i < doc.page_count]
        if not indexes:
            raise ConversionError("No pages to convert.")
        self.report.pages = len(indexes)
        scanned = {i for i in indexes if is_scanned(doc[i])}
        typed = [i for i in indexes if i not in scanned]

        parsed = self._parse_typed(doc, typed) if typed else {}
        ocr_results = self._ocr_scanned(doc, [i for i in indexes if i in scanned]) if scanned else {}

        self.progress(0.95, "Writing the Word document...")
        word = Document()
        for i in indexes:
            self._check_cancel()
            if i in parsed and self._write_parsed(word, parsed[i]):
                self.report.converted.append(i + 1)
            elif i in ocr_results:
                self._write_blocks(word, doc[i], ocr_results[i])
                self.report.ocr.append(i + 1)
            else:
                self._write_picture(word, doc[i])
                self.report.pictures.append(i + 1)
        self._save(word, out_path)
        self.progress(1.0, "Done")
        return self.report

    def _save(self, word, out_path):
        folder = os.path.dirname(os.path.abspath(out_path))
        fd, tmp = tempfile.mkstemp(suffix=".docx", dir=folder)
        os.close(fd)
        try:
            word.save(tmp)
            os.replace(tmp, out_path)
        except PermissionError as e:
            raise ConversionError(f"Could not write {out_path}. If it is open in Word, close it and try again.") from e
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)

    # ------------------------------------------------------------ typed pages
    def _parse_typed(self, doc, typed) -> dict:
        """{page index: pdf2docx Page} for every page pdf2docx could parse."""
        from pdf2docx import Converter

        logging.getLogger("pdf2docx").setLevel(logging.ERROR)
        data = doc.tobytes()
        cv = Converter(stream=data)
        settings = cv.default_settings
        settings.update(raw_exceptions=True, ignore_page_error=False, multi_processing=False, debug=False)
        self.progress(0.02, "Analysing the document layout...")
        try:
            cv.load_pages(pages=typed).parse_document(**settings)
            to_parse = [(i, cv.pages[i]) for i in typed]
        except Exception as e:  # noqa: BLE001 - retry the pages one by one below
            logging.warning("pdf2docx document analysis failed: %s", e)
            to_parse = None
        parsed = {}
        if to_parse is None:  # analyse each page on its own, so one bad page can't sink the rest
            for n, i in enumerate(typed):
                self._check_cancel()
                self.progress(0.05 + 0.55 * n / len(typed), f"Converting page {i + 1}...")
                try:
                    single = Converter(stream=data)
                    single.load_pages(pages=[i]).parse_document(**settings)
                    single.pages[i].parse(**settings)
                    parsed[i] = single.pages[i]
                except Exception as e:  # noqa: BLE001
                    self.report.warnings.append(f"Page {i + 1} could not be converted to text ({e}).")
            return parsed
        for n, (i, page) in enumerate(to_parse):
            self._check_cancel()
            self.progress(0.05 + 0.55 * n / len(to_parse), f"Converting page {i + 1} ({n + 1} of {len(to_parse)})...")
            try:
                page.parse(**settings)
                parsed[i] = page
            except Exception as e:  # noqa: BLE001
                self.report.warnings.append(f"Page {i + 1} could not be converted to text ({e}).")
        return parsed

    def _write_parsed(self, word, page) -> bool:
        body = word.element.body
        before = len(body)
        try:
            page.make_docx(word)
        except Exception as e:  # noqa: BLE001 - drop the half-written page, use a picture instead
            for child in list(body)[before:]:
                if not child.tag.endswith("}sectPr"):
                    body.remove(child)
            self.report.warnings.append(f"Page {page.id + 1} could not be laid out in Word ({e}).")
            return False
        return True

    # ------------------------------------------------------------ scanned pages
    def _pick_engine(self):
        if self.ocr == OCR_NONE:
            return None
        if self.ocr in (OCR_AUTO, OCR_TESSERACT) and tesseract_available():
            return OCR_TESSERACT
        if self.ocr == OCR_TESSERACT:
            self.report.warnings.append("Tesseract OCR is not installed, so the GOT-OCR 2.0 model was used.")
        return OCR_TROCR if self.ocr == OCR_TROCR else OCR_GOT

    def _ocr_scanned(self, doc, scanned) -> dict:
        engine = self._pick_engine()
        if engine is None:
            return {}
        results = {}
        if engine == OCR_TESSERACT:
            self.report.ocr_engine = engine
            for n, i in enumerate(scanned):
                self._check_cancel()
                self.progress(0.6 + 0.35 * n / len(scanned), f"Reading scanned page {i + 1} (Tesseract)...")
                try:
                    results[i] = tesseract_blocks(doc[i])
                except Exception as e:  # noqa: BLE001
                    self.report.warnings.append(f"OCR failed on page {i + 1} ({e}).")
            return results
        return self._model_ocr(doc, scanned, engine)

    def _model_ocr(self, doc, scanned, engine) -> dict:
        from . import ocr_models

        label = ocr_models.MODEL_LABELS[engine]
        current = {"text": ""}

        def status(text):
            self.progress(current.get("fraction", 0.6), f"{current['text']} {text}".strip())

        self.progress(0.6, f"Starting the {label} model (the first start can take a minute)...")
        try:
            self._model = ocr_models.ModelClient(engine, status=status)
        except ocr_models.ModelError as e:
            self._check_cancel()
            self.report.warnings.append(f"The {label} model could not be started, so scanned pages were kept "
                                        f"as pictures. {e}")
            return {}
        self.report.ocr_engine = engine
        results = {}
        try:
            with tempfile.TemporaryDirectory() as tmp:
                for n, i in enumerate(scanned):
                    self._check_cancel()
                    current.update(fraction=0.6 + 0.35 * n / len(scanned),
                                   text=f"Scanned page {i + 1} ({n + 1} of {len(scanned)}):")
                    status(f"reading with {label}...")
                    image = os.path.join(tmp, f"page{i + 1}.png")
                    doc[i].get_pixmap(dpi=200 if engine == OCR_TROCR else 144).save(image)
                    try:
                        text = self._model.read_page(image)
                    except ocr_models.ModelError as e:
                        self._check_cancel()
                        self.report.warnings.append(f"OCR failed on page {i + 1} ({e}).")
                        if self._model.proc.poll() is not None:
                            break  # the worker died; the rest stay pictures
                        continue
                    results[i] = self._model_result(engine, text)
        finally:
            self._model.close()
            self._model = None
        return results

    def _model_result(self, engine, text):
        """A model's page text -> what _ocr_scanned returns for that page."""
        return got_blocks(text) if engine == OCR_GOT else trocr_blocks(text)

    # ------------------------------------------------------------ writing
    def _new_page(self, word, page, margin):
        from docx.enum.section import WD_SECTION
        from docx.shared import Pt

        section = word.sections[0] if is_empty(word) else word.add_section(WD_SECTION.NEW_PAGE)
        section.page_width = Pt(page.rect.width)
        section.page_height = Pt(page.rect.height)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Pt(margin))
        return section

    def _write_blocks(self, word, page, blocks):
        self._new_page(word, page, OCR_MARGIN_PT)
        if not blocks:
            word.add_paragraph("")
        for block in blocks:
            if block[0] == "heading":
                word.add_heading(block[2], level=min(block[1], 9))
            elif block[0] == "table":
                rows = block[1]
                cols = max(len(r) for r in rows)
                table = word.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        table.cell(r, c).text = cell
                word.add_paragraph("")
            else:
                word.add_paragraph(block[1])
        if self.keep_scan_images:
            self._add_page_picture(word, page, OCR_MARGIN_PT)

    def _write_picture(self, word, page):
        self._new_page(word, page, PICTURE_MARGIN_PT)
        self._add_page_picture(word, page, PICTURE_MARGIN_PT)

    @staticmethod
    def _add_page_picture(word, page, margin):
        from docx.shared import Pt

        pix = page.get_pixmap(dpi=PICTURE_DPI)
        # Word needs a little slack below the picture or it spills onto a new page
        avail_w = page.rect.width - 2 * margin
        avail_h = page.rect.height - 2 * margin - 14
        scale = min(avail_w / page.rect.width, avail_h / page.rect.height)
        word.add_picture(io.BytesIO(pix.tobytes("png")), width=Pt(page.rect.width * scale))


def is_empty(word) -> bool:
    """True while the document is still empty, so the first page reuses the
    section every new Word document starts with (as pdf2docx does)."""
    body = word.element.body
    return all(child.tag.endswith("}sectPr") for child in body)


def convert(source, out_path, **options) -> ConversionReport:
    """One-call form of PdfToDocx(source, **options).run(out_path)."""
    return PdfToDocx(source, **options).run(out_path)
