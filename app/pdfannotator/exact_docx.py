"""Exact-layout Word: every page of the PDF becomes a Word page that looks
the way it did.

- The page's graphics (rules, table borders, charts, pictures) are a
  picture behind the text, made from the page with its text taken out.
- Every line of text is a Word frame (a paragraph placed at an exact spot on
  the page) holding real, editable text in a matching font: Arial for
  Helvetica / Nimbus Sans, Times New Roman for Times, Courier New for
  Courier, Cambria for Computer Modern, Cambria Math for maths, or the
  PDF's own font when Windows has it. Each word is placed where it was: the
  spacing after it is adjusted (measured with the font's own widths) so the
  next one starts at its original position.
- Glyphs that can't be retyped (display integrals, odd symbol fonts) are
  small pictures at their spot.

Word puts a line's baseline one font size below the top of a frame whose
line height is 1.25 times the size (measured from Word's own PDF export);
the frames are placed with that.

No Qt here."""
import io
import os
from functools import lru_cache

import pymupdf as fitz

from . import exact_layout as ex

BACKGROUND_DPI = 220
GLYPH_DPI = 400
LINE_FACTOR = 1.25          # frame line height, in font sizes
EMU_PER_PT = 12700

_FONTS_DIR = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
# family -> (Word font, {(bold, italic): (file, index in a collection)})
_WORD_FONTS = {
    ex.SANS: ("Arial", {(0, 0): ("arial.ttf", 0), (1, 0): ("arialbd.ttf", 0), (0, 1): ("ariali.ttf", 0),
                        (1, 1): ("arialbi.ttf", 0)}),
    ex.SERIF: ("Times New Roman", {(0, 0): ("times.ttf", 0), (1, 0): ("timesbd.ttf", 0), (0, 1): ("timesi.ttf", 0),
                                   (1, 1): ("timesbi.ttf", 0)}),
    ex.MONO: ("Courier New", {(0, 0): ("cour.ttf", 0), (1, 0): ("courbd.ttf", 0), (0, 1): ("couri.ttf", 0),
                              (1, 1): ("courbi.ttf", 0)}),
    ex.CM_SERIF: ("Cambria", {(0, 0): ("cambria.ttc", 0), (1, 0): ("cambriab.ttf", 0), (0, 1): ("cambriai.ttf", 0),
                              (1, 1): ("cambriaz.ttf", 0)}),
    ex.CM_SANS: ("Arial", {(0, 0): ("arial.ttf", 0), (1, 0): ("arialbd.ttf", 0), (0, 1): ("ariali.ttf", 0),
                           (1, 1): ("arialbi.ttf", 0)}),
    ex.CM_MONO: ("Courier New", {(0, 0): ("cour.ttf", 0), (1, 0): ("courbd.ttf", 0), (0, 1): ("couri.ttf", 0),
                                 (1, 1): ("courbi.ttf", 0)}),
    ex.MATH: ("Cambria Math", {(0, 0): ("cambria.ttc", 1)}),
}


@lru_cache(maxsize=64)
def _metrics(path, index):
    """(advance widths by character, units per em) of a font file, or None."""
    try:
        from fontTools.ttLib import TTCollection, TTFont

        if path.lower().endswith((".ttc", ".otc")):
            font = TTCollection(path, lazy=True).fonts[index]
        else:
            font = TTFont(path, lazy=True, fontNumber=index)
        cmap = font.getBestCmap() or {}
        hmtx = font["hmtx"].metrics
        upm = font["head"].unitsPerEm
        widths = {chr(code): hmtx[glyph][0] for code, glyph in cmap.items() if glyph in hmtx}
        return widths, upm
    except Exception:  # noqa: BLE001 - no correction for this font
        return None


class FontMap:
    """The Word font for a segment, and how wide its text is in that font."""

    def __init__(self):
        self._system = {}

    def word_font(self, seg):
        if seg.system_font:
            found = self._system_font(seg.system_font)
            if found:
                return found
        name, files = _WORD_FONTS.get(seg.family, _WORD_FONTS[ex.SANS])
        bold, italic = (0, 0) if seg.family == ex.MATH else (int(seg.bold), int(seg.italic))
        file, index = files.get((bold, italic)) or files[(0, 0)]
        return name, os.path.join(_FONTS_DIR, file), index

    def _system_font(self, family):
        if family not in self._system:
            from . import fonts

            files = None
            try:
                files = fonts.family_files(family)
            except Exception:  # noqa: BLE001
                files = None
            path = files.get("regular") if files else None
            ok = bool(path) and str(path).lower().startswith(_FONTS_DIR.lower())
            self._system[family] = (family, str(path), 0) if ok else None
        return self._system[family]

    def width(self, text, seg):
        """Width in points of `text` in the segment's Word font (None: unknown)."""
        _name, path, index = self.word_font(seg)
        metrics = _metrics(path, index) if path and os.path.isfile(path) else None
        if metrics is None:
            return None
        widths, upm = metrics
        size = _word_size(seg.size)
        fallback = widths.get("n", upm // 2)
        return sum(widths.get(ch, fallback) for ch in text) * size / upm


def _word_size(size):
    """Word sizes are in half points."""
    return max(1.0, round(size * 2) / 2)


def _q(tag):
    from docx.oxml.ns import qn

    return qn(tag)


def _el(tag, **attrs):
    from docx.oxml import OxmlElement

    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(_q(f"w:{k}") if ":" not in k else _q(k), str(v))
    return e


class ExactDocx:
    """Writes PDF pages into a python-docx Document with the exact layout."""

    def __init__(self, word):
        self.word = word
        self.fonts = FontMap()
        self._pictures = 0

    # ---- pages
    def new_page(self, width, height):
        from docx.enum.section import WD_SECTION
        from docx.shared import Pt

        body = self.word.element.body
        first = all(child.tag.endswith("}sectPr") for child in body)
        section = self.word.sections[0] if first else self.word.add_section(WD_SECTION.NEW_PAGE)
        section.page_width, section.page_height = Pt(width), Pt(height)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance",
                     "footer_distance", "gutter"):
            setattr(section, side, Pt(0))
        # the paragraph that holds the page's background (and keeps the page)
        para = self.word.add_paragraph()
        self._tight(para, 1)
        return para

    @staticmethod
    def _tight(para, size_pt):
        ppr = para._p.get_or_add_pPr()
        ppr.append(_el("w:spacing", before=0, after=0, line=max(20, int(size_pt * 20)), lineRule="exact"))
        rpr = _el("w:rPr")
        rpr.append(_el("w:sz", val=2))
        ppr.append(rpr)

    def picture(self, para, png, x, y, w, h, behind):
        """A picture placed on the page at (x, y) (top-left, points), behind
        the text (the background) or in front of it."""
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        from docx.shared import Pt

        run = para.add_run()
        run.add_picture(io.BytesIO(png), width=Pt(w), height=Pt(h))
        inline = run._r.find(".//" + _q("wp:inline"))
        self._pictures += 1
        anchor = parse_xml(
            f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{self._pictures}" behindDoc="{1 if behind else 0}" locked="0" layoutInCell="1" '
            f'allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{int(round(x * EMU_PER_PT))}</wp:posOffset>'
            f'</wp:positionH><wp:positionV relativeFrom="page"><wp:posOffset>{int(round(y * EMU_PER_PT))}'
            f'</wp:posOffset></wp:positionV></wp:anchor>')
        anchor.append(inline.find(_q("wp:extent")))
        anchor.append(parse_xml(f'<wp:effectExtent {nsdecls("wp")} l="0" t="0" r="0" b="0"/>'))
        anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
        for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
            child = inline.find(_q(tag))
            if child is not None:
                anchor.append(child)
        inline.getparent().replace(inline, anchor)

    # ---- lines
    def line(self, line, page_w):
        """One line of text as a frame paragraph at its exact place."""
        segs = sorted((s for s in line.segments if not s.picture and s.text.strip()), key=lambda s: s.x)
        if not segs:
            return
        main = max(segs, key=lambda s: (s.size, -abs(s.y - segs[0].y)))
        size = _word_size(main.size)
        height = LINE_FACTOR * size
        top = main.y - size                       # where Word's baseline lands
        x0 = min(s.x for s in segs)
        right = max(s.x + s.width for s in segs)
        para = self.word.add_paragraph()
        ppr = para._p.get_or_add_pPr()
        frame_w = min(page_w - x0, (right - x0) * 1.15 + 24)
        ppr.append(_el("w:framePr", w=int(max(frame_w, 12) * 20), h=int(height * 20), hRule="exact", wrap="around",
                       vAnchor="page", hAnchor="page", x=int(round(x0 * 20)), y=int(round(top * 20))))
        ppr.append(_el("w:spacing", before=0, after=0, line=int(round(height * 20)), lineRule="exact"))
        ppr.append(_el("w:ind", left=0, right=0, firstLine=0))
        ppr.append(_el("w:jc", val="left"))
        pieces = []                               # (text, segment, x where it should start)
        for seg in segs:
            words = seg.words or [(seg.text, seg.x)]
            for k, (word, wx) in enumerate(words):
                pieces.append((word, seg, wx, k < len(words) - 1))
        for n, (word, seg, wx, space_after) in enumerate(pieces):
            nxt = pieces[n + 1][2] if n + 1 < len(pieces) else None
            text = word + (" " if space_after else "")
            natural = self.fonts.width(text, seg)
            if nxt is None or natural is None:
                self._run(para, text, seg, main, 0)
                continue
            delta = (nxt - wx) - natural          # what the gap after this word must add
            if space_after:
                self._run(para, word, seg, main, 0)
                self._run(para, " ", seg, main, delta)
            elif len(word) > 1:
                self._run(para, word[:-1], seg, main, 0)
                self._run(para, word[-1], seg, main, delta)
            else:
                self._run(para, word, seg, main, delta)

    def _run(self, para, text, seg, main, spacing_pt):
        from docx.shared import Pt, RGBColor

        if not text:
            return
        run = para.add_run(text)
        name, _path, _index = self.fonts.word_font(seg)
        font = run.font
        font.name = name
        rpr = run._r.get_or_add_rPr()
        fonts_el = rpr.find(_q("w:rFonts"))
        if fonts_el is not None:
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts_el.set(_q(attr), name)
        font.size = Pt(_word_size(seg.size))
        if seg.bold and seg.family != ex.MATH:
            font.bold = True
        if seg.italic and seg.family != ex.MATH:
            font.italic = True
        if seg.color != (0, 0, 0):
            font.color.rgb = RGBColor(*seg.color)
        twips = int(round(spacing_pt * 20))
        if twips:
            rpr.append(_el("w:spacing", val=max(-1584, min(1584, twips))))
        shift = round((main.y - seg.y) * 2)     # half points up (superscript) or down (subscript)
        if shift:
            rpr.append(_el("w:position", val=shift))

    # ---- a whole page
    def page(self, page, layout, background_png):
        para = self.new_page(layout.width, layout.height)
        if background_png:
            self.picture(para, background_png, 0, 0, layout.width, layout.height, behind=True)
        glyphs = ex.glyph_page(page, layout)
        try:
            for line in layout.lines:
                if line.angle:
                    self._turned(para, line, page, glyphs)
                    continue
                self.line(line, layout.width)
                for seg in line.segments:
                    if seg.picture:
                        self._glyph(para, seg, page, glyphs)
        finally:
            if glyphs:
                glyphs[0].close()

    def _glyph(self, para, seg, page, glyphs):
        rect = ex.picture_rect(seg, page.rect)
        if rect.is_empty or rect.width < 0.5 or rect.height < 0.5:
            return
        source = glyphs[1] if glyphs else page
        pix = source.get_pixmap(dpi=GLYPH_DPI, clip=rect * ~page.rotation_matrix, alpha=True)
        self.picture(para, pix.tobytes("png"), rect.x0, rect.y0, rect.width, rect.height, behind=False)

    def _turned(self, para, line, page, glyphs):
        """Sideways text: kept as a picture of itself (frames can't turn)."""
        rect = fitz.Rect()
        for seg in line.segments:
            rect |= fitz.Rect(seg.bbox)
        rect &= page.rect
        if rect.is_empty:
            return
        pix = page.get_pixmap(dpi=GLYPH_DPI, clip=rect * ~page.rotation_matrix, alpha=True)
        self.picture(para, pix.tobytes("png"), rect.x0, rect.y0, rect.width, rect.height, behind=False)


def background_png(bg_page, dpi=BACKGROUND_DPI):
    """A text-free page as a PNG (None when it is blank)."""
    import numpy as np

    pix = bg_page.get_pixmap(dpi=dpi, alpha=False)
    samples = np.frombuffer(pix.samples, dtype=np.uint8)
    if samples.size == 0 or samples.min() == samples.max():
        return None
    return pix.tobytes("png")
